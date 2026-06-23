import os
import re
import csv
import io
import time
import asyncio
import json
import joblib
import psycopg2
import requests
from datetime import datetime, timezone, timedelta
from fastapi import FastAPI, HTTPException, Query, Request, Header, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, Response
from pydantic import BaseModel
from typing import Optional
from collections import defaultdict
from jose import jwt, JWTError

from src.fabric_client import FabricClient
from src.private_store import PrivateRecordStore
from src.kyc_store import KYCStore
from src.user_store import UserStore, verify_password

app = FastAPI(title="Fraud Detection API")

# Allowed browser origins for the hosted frontend.
#  - myfaid.com / www.myfaid.com : the Vercel production domain
#  - localhost                   : local dev (frontend on :3000)
#  - *.vercel.app                : Vercel preview/initial deploy URLs
# Extra origins can be added via CORS_EXTRA_ORIGINS (comma-separated).
_cors_origins = [
    "https://myfaid.com",
    "https://www.myfaid.com",
    "http://localhost:3000",
    "http://127.0.0.1:3000",
]
_cors_origins += [o.strip() for o in os.getenv("CORS_EXTRA_ORIGINS", "").split(",") if o.strip()]

app.add_middleware(
    CORSMiddleware,
    allow_origins=_cors_origins,
    allow_origin_regex=r"https://.*\.vercel\.app",
    allow_methods=["*"],
    allow_headers=["*"],
)

FABRIC_GATEWAY_URL = os.getenv("FABRIC_GATEWAY_URL", "http://localhost:8080")
POSTGRES_DSN       = os.getenv("POSTGRES_DSN", "postgresql://fraud_user:fraud_pass@localhost:5432/fraud_private")
ENCRYPTION_KEY     = os.getenv("ENCRYPTION_KEY", "").encode()
JWT_SECRET         = os.getenv("JWT_SECRET", "fraud-detection-secret-change-in-prod")
JWT_ALGORITHM      = "HS256"
JWT_EXPIRE_HOURS   = 8
# Google "Sign in with Google" (SSO). Same Client ID as the frontend button.
# Public OAuth Client ID (safe to embed); env var overrides it. Used only as the
# expected audience when verifying Google ID tokens, so a default is safe.
GOOGLE_CLIENT_ID   = os.getenv("GOOGLE_CLIENT_ID", "315587208584-cuu3e21ppm52t1a4bioqv4aedp2glpmj.apps.googleusercontent.com")

fabric = FabricClient(gateway_url=FABRIC_GATEWAY_URL)

# ── SSE client registry ───────────────────────────────────────
# One asyncio.Queue per connected browser tab.
# When data changes, we put an event into every queue.
_sse_clients: set[asyncio.Queue] = set()
_event_loop: asyncio.AbstractEventLoop | None = None


@app.on_event("startup")
async def _capture_event_loop():
    global _event_loop
    _event_loop = asyncio.get_event_loop()


def _notify_clients(event_type: str, data: dict):
    """
    Notify every active SSE client. Safe to call from both async routes
    and sync routes (which FastAPI runs in a thread-pool thread).
    asyncio.Queue is not thread-safe, so we always schedule via
    call_soon_threadsafe to ensure the put runs on the event loop thread.
    """
    if not _sse_clients or _event_loop is None:
        return
    payload = {"type": event_type, **data}
    for q in list(_sse_clients):
        try:
            _event_loop.call_soon_threadsafe(q.put_nowait, payload)
        except Exception:
            pass


def get_db():
    return psycopg2.connect(POSTGRES_DSN)


# ── Dashboard snapshot builder ────────────────────────────────
# Reads from PostgreSQL (the read model) so the snapshot is available
# immediately after the fraud consumer saves — no need to wait for
# blockchain confirmation. Runs in a thread executor so it doesn't
# block the event loop.

def _build_snapshot() -> dict:
    conn = get_db()
    try:
        with conn.cursor() as cur:
            # ── Stats ─────────────────────────────────────────────
            cur.execute("""
                SELECT
                    COUNT(*)                                                            AS total,
                    COUNT(*) FILTER (WHERE predicted_label = 'FRAUD')                  AS fraud_count,
                    COUNT(*) FILTER (WHERE predicted_label = 'LEGIT')                  AS legit_count,
                    COUNT(*) FILTER (WHERE predicted_label = 'FRAUD'
                                     AND reviewed_by IS NULL)                          AS pending_review,
                    COUNT(*) FILTER (WHERE predicted_label = 'FRAUD'
                                     AND reviewed_by IS NOT NULL)                      AS reviewed,
                    COALESCE(AVG(fraud_score), 0)                                      AS avg_fraud_score,
                    COALESCE(SUM(amount_myr) FILTER (WHERE predicted_label = 'FRAUD'), 0) AS balance_at_risk
                FROM private_transactions
            """)
            total, fraud_count, legit_count, pending, reviewed, avg_score, balance_at_risk = cur.fetchone()
            cur.execute("SELECT COUNT(*) FROM frozen_cards")
            compromised_cards = cur.fetchone()[0]

            # ── Trend (monthly) ───────────────────────────────────
            cur.execute("""
                SELECT TO_CHAR(DATE_TRUNC('month', timestamp), 'YYYY-MM') AS month,
                       COUNT(*) FILTER (WHERE predicted_label = 'FRAUD')                     AS fraud,
                       COUNT(*) FILTER (WHERE predicted_label = 'LEGIT')                     AS legit,
                       COALESCE(SUM(amount_myr) FILTER (WHERE predicted_label = 'FRAUD'), 0) AS amount_at_risk
                FROM private_transactions
                WHERE predicted_label IN ('FRAUD', 'LEGIT')
                GROUP BY DATE_TRUNC('month', timestamp)
                ORDER BY 1
            """)
            trend_list = [
                {"date": m, "fraud": int(f), "legit": int(l), "amount_at_risk": float(a)}
                for m, f, l, a in cur.fetchall()
            ]

            # ── Score distribution (FRAUD only, 3 buckets) ────────
            cur.execute("""
                SELECT
                    COUNT(*) FILTER (WHERE fraud_score < 0.6)                        AS low,
                    COUNT(*) FILTER (WHERE fraud_score >= 0.6 AND fraud_score < 0.8) AS medium,
                    COUNT(*) FILTER (WHERE fraud_score >= 0.8)                       AS high
                FROM private_transactions
                WHERE predicted_label = 'FRAUD' AND fraud_score IS NOT NULL
            """)
            low, medium, high = cur.fetchone()

            # ── Recent transactions ────────────────────────────────
            cur.execute("""
                SELECT
                    transaction_id, timestamp::text, amount_myr, merchant_name,
                    customer_ref, fraud_score, predicted_label, ml_prediction,
                    rule_flag, risk_reasons, ground_truth_label,
                    reviewed_by, reviewed_at::text,
                    mcc, mode, location, ic_hash, masked_card_number,
                    EXISTS (SELECT 1 FROM frozen_cards fc
                            WHERE fc.card_hash = private_transactions.card_hash) AS card_frozen
                FROM private_transactions
                ORDER BY timestamp DESC
                LIMIT 8
            """)
            cols = [
                "transaction_id", "timestamp", "amount_myr", "merchant_name",
                "customer_ref", "fraud_score", "predicted_label", "ml_prediction",
                "rule_flag", "risk_reasons", "ground_truth_label",
                "reviewed_by", "reviewed_at",
                "mcc", "mode", "location", "ic_hash", "masked_card_number",
                "card_frozen",
            ]
            recent = []
            for r in cur.fetchall():
                txn = dict(zip(cols, r))
                txn["fraud_score"] = float(txn["fraud_score"]) if txn["fraud_score"] is not None else 0.0
                txn["amount_myr"]  = float(txn["amount_myr"])  if txn["amount_myr"]  is not None else 0.0
                rr = txn["risk_reasons"]
                txn["risk_reasons"] = json.loads(rr) if isinstance(rr, str) else (rr or [])
                recent.append(txn)
    finally:
        conn.close()

    return {
        "stats": {
            "total":             int(total),
            "fraud_count":       int(fraud_count),
            "legit_count":       int(legit_count),
            "pending_review":    int(pending),
            "reviewed":          int(reviewed),
            "avg_fraud_score":   round(float(avg_score), 4),
            "balance_at_risk":   float(balance_at_risk),
            "compromised_cards": int(compromised_cards),
        },
        "trend": trend_list,
        "score_distribution": [
            {"range": "Low",    "count": int(low)},
            {"range": "Medium", "count": int(medium)},
            {"range": "High",   "count": int(high)},
        ],
        "recent_transactions": recent,
    }


# ── Models ────────────────────────────────────────────────────

class LoginRequest(BaseModel):
    username: str
    password: str


class GoogleLoginRequest(BaseModel):
    # The signed ID token the Google "Sign in" button hands back in the browser.
    credential: str


class ChangePasswordRequest(BaseModel):
    current_password: str
    new_password: str


class ReviewRequest(BaseModel):
    ground_truth_label: int
    reviewer_id: str
    notes: Optional[str] = ""


class IotCartRequest(BaseModel):
    # The "heist checkout" shop page arms this; the Arduino bridge reads it on a
    # physical card tap and turns it into one real transaction. Semantic fields
    # only — the bridge maps foreign/online to IP/device/mode.
    label:    str                       # display name, e.g. "Luxury Watch"
    amount:   float
    merchant: str
    mcc:      str
    location: str
    online:   bool = True               # True -> Mode "Online", else "In-Person"
    foreign:  bool = True               # True -> foreign IP (cross-location)


# ── Auth ──────────────────────────────────────────────────────

@app.post("/api/auth/login")
def login(body: LoginRequest):
    conn = get_db()
    try:
        user = UserStore(conn).get_by_username(body.username)
    finally:
        conn.close()

    if not user or not verify_password(body.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid username or password")

    return _issue_token(user)


def _issue_token(user: dict) -> dict:
    """Mint the app's own JWT for a resolved user row. Both password login and
    Google SSO funnel through here, so everything downstream (roles, /me, review
    attribution) is identical regardless of how the user authenticated."""
    payload = {
        "sub":     user["username"],
        "user_id": user["user_id"],
        "role":    user["role"],
        "exp":     datetime.now(timezone.utc) + timedelta(hours=JWT_EXPIRE_HOURS),
    }
    token = jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)
    return {
        "token": token,
        "user":  {"user_id": user["user_id"], "username": user["username"], "role": user["role"]},
    }


@app.post("/api/auth/google")
def google_login(body: GoogleLoginRequest):
    """Sign in with Google (SSO). The browser button returns a signed Google ID
    token; we verify it with Google, then map the verified email to a local user
    (email column = allowlist). No matching email → not authorized. Either way the
    user ends up with the same app JWT as a password login."""
    if not GOOGLE_CLIENT_ID:
        raise HTTPException(status_code=503, detail="Google sign-in is not configured on the server.")

    # Verify signature, audience (our Client ID), issuer and expiry against Google.
    try:
        from google.oauth2 import id_token as google_id_token
        from google.auth.transport import requests as google_requests
        info = google_id_token.verify_oauth2_token(
            body.credential, google_requests.Request(), GOOGLE_CLIENT_ID
        )
    except Exception:
        raise HTTPException(status_code=401, detail="Could not verify your Google sign-in.")

    email = (info.get("email") or "").strip().lower()
    if not email or not info.get("email_verified"):
        raise HTTPException(status_code=401, detail="Your Google account has no verified email.")

    conn = get_db()
    try:
        user = UserStore(conn).get_by_email(email)
    finally:
        conn.close()

    if not user:
        raise HTTPException(
            status_code=403,
            detail=f"{email} isn't authorized. Ask an admin to add your email to an account.",
        )
    return _issue_token(user)


@app.get("/api/auth/me")
def get_me(authorization: Optional[str] = Header(None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(authorization[7:], JWT_SECRET, algorithms=[JWT_ALGORITHM])
        return {
            "user_id":  payload.get("user_id"),
            "username": payload.get("sub"),
            "role":     payload.get("role"),
        }
    except JWTError:
        raise HTTPException(status_code=401, detail="Invalid or expired token")


def _require_username(authorization: Optional[str]) -> str:
    """Decode the Bearer token and return the username (sub), or 401."""
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(authorization[7:], JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except JWTError:
        raise HTTPException(status_code=401, detail="Invalid or expired token")
    username = payload.get("sub")
    if not username:
        raise HTTPException(status_code=401, detail="Invalid token")
    return username


@app.get("/api/auth/me/review-stats")
def my_review_stats(authorization: Optional[str] = Header(None)):
    """Review activity for the logged-in analyst (keyed on reviewed_by = their
    username). Honest, per-user metrics — follows the account, not a hardcoded name."""
    username = _require_username(authorization)
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT
                    COUNT(*)                                                    AS reviews_done,
                    COUNT(*) FILTER (WHERE ground_truth_label = 1)              AS confirmed_fraud,
                    COUNT(*) FILTER (WHERE ground_truth_label = 0)              AS cleared,
                    COALESCE(SUM(amount_myr) FILTER (WHERE ground_truth_label = 0), 0) AS amount_approved
                FROM private_transactions
                WHERE reviewed_by = %s AND reviewed_by <> ''
                """,
                (username,),
            )
            reviews_done, confirmed_fraud, cleared, amount_approved = cur.fetchone()

            # Cards this analyst cleared (marked legitimate) that are no longer frozen.
            cur.execute(
                """
                SELECT COUNT(DISTINCT pt.card_hash)
                FROM private_transactions pt
                WHERE pt.reviewed_by = %s
                  AND pt.ground_truth_label = 0
                  AND pt.card_hash IS NOT NULL
                  AND NOT EXISTS (
                      SELECT 1 FROM frozen_cards f WHERE f.card_hash = pt.card_hash
                  )
                """,
                (username,),
            )
            cards_unfrozen = cur.fetchone()[0]
    finally:
        conn.close()

    return {
        "reviews_done":    int(reviews_done),
        "confirmed_fraud": int(confirmed_fraud),
        "cleared":         int(cleared),
        "amount_approved": float(amount_approved or 0),
        "cards_unfrozen":  int(cards_unfrozen or 0),
    }


@app.post("/api/auth/change-password")
def change_password(body: ChangePasswordRequest, authorization: Optional[str] = Header(None)):
    username = _require_username(authorization)
    if len(body.new_password) < 6:
        raise HTTPException(status_code=400, detail="New password must be at least 6 characters")
    if body.new_password == body.current_password:
        raise HTTPException(status_code=400, detail="New password must differ from the current one")

    conn = get_db()
    try:
        store = UserStore(conn)
        user = store.get_by_username(username)
        if not user or not verify_password(body.current_password, user["password_hash"]):
            raise HTTPException(status_code=401, detail="Current password is incorrect")
        store.update_password(username, body.new_password)
    finally:
        conn.close()
    return {"status": "SUCCESS"}


# ── Transactions ──────────────────────────────────────────────

@app.get("/api/transactions")
def list_transactions(
    decision: Optional[str] = Query(None),
    reviewed: Optional[bool] = Query(None),
    search: Optional[str] = Query(None),
    risk: Optional[str] = Query(None),
    limit: int = Query(50, le=100000),
    offset: int = Query(0),
):
    conditions, params = ["1=1"], []

    if decision:
        conditions.append("predicted_label = %s")
        params.append(decision.upper())
    if reviewed is True:
        conditions.append("reviewed_by IS NOT NULL AND reviewed_by != ''")
    elif reviewed is False:
        conditions.append("(reviewed_by IS NULL OR reviewed_by = '')")
    if search:
        conditions.append(
            "(transaction_id ILIKE %s OR merchant_name ILIKE %s OR customer_ref ILIKE %s)"
        )
        s = f"%{search}%"
        params.extend([s, s, s])
    if risk:
        r = risk.upper()
        if r == "LOW":
            conditions.append("fraud_score < 0.6")
        elif r == "MEDIUM":
            conditions.append("fraud_score >= 0.6 AND fraud_score < 0.8")
        elif r == "HIGH":
            conditions.append("fraud_score >= 0.8")

    where = " AND ".join(conditions)
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute(f"SELECT COUNT(*) FROM private_transactions WHERE {where}", params)
            total = cur.fetchone()[0]
            cur.execute(
                f"""
                SELECT transaction_id, customer_ref, timestamp, amount_myr,
                       merchant_name, mcc, mode, location, ic_hash,
                       masked_card_number, fraud_score, predicted_label,
                       ml_prediction, rule_flag, risk_reasons,
                       ground_truth_label, reviewed_by, reviewed_at, created_at,
                       (SELECT fc.frozen_at FROM frozen_cards fc
                          WHERE fc.card_hash = private_transactions.card_hash) AS card_frozen_at,
                       (SELECT fc.trigger_reasons FROM frozen_cards fc
                          WHERE fc.card_hash = private_transactions.card_hash) AS card_frozen_reasons
                FROM private_transactions
                WHERE {where}
                ORDER BY timestamp DESC
                LIMIT %s OFFSET %s
                """,
                params + [limit, offset],
            )
            rows = cur.fetchall()
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    def _row_to_txn(row):
        try:
            risk_reasons = json.loads(row[14]) if row[14] else []
        except Exception:
            risk_reasons = []
        frozen_at = row[19]
        try:
            frozen_reasons = json.loads(row[20]) if row[20] else []
        except Exception:
            frozen_reasons = []
        return {
            "transaction_id":     row[0],
            "customer_ref":       row[1] or "",
            "timestamp":          row[2].isoformat() if row[2] else "",
            "amount_myr":         float(row[3]) if row[3] is not None else 0.0,
            "merchant_name":      row[4] or "",
            "mcc":                row[5] or "",
            "mode":               row[6] or "",
            "location":           row[7] or "",
            "ic_hash":            row[8] or "",
            "masked_card_number": row[9] or "",
            "fraud_score":        float(row[10]) if row[10] is not None else 0.0,
            "predicted_label":    row[11] or "pending",
            "ml_prediction":      int(row[12]) if row[12] is not None else 0,
            "rule_flag":          int(row[13]) if row[13] is not None else 0,
            "risk_reasons":       risk_reasons,
            "ground_truth_label": int(row[15]) if row[15] is not None else 0,
            "reviewed_by":        row[16] or "",
            "reviewed_at":        row[17].isoformat() if row[17] else "",
            "created_at":         row[18].isoformat() if row[18] else "",
            "created_by":         "",
            "card_frozen":         frozen_at is not None,
            "card_frozen_at":      frozen_at.isoformat() if frozen_at else "",
            "card_frozen_reasons": frozen_reasons,
        }

    return {"total": int(total), "transactions": [_row_to_txn(r) for r in rows]}


# ── IoT "armed cart" (booth shop page <-> Arduino bridge) ─────────────────────
# A single in-memory slot. The /shop page POSTs what the visitor wants to buy;
# the Arduino bridge GETs it on a physical card tap, fires the transaction, then
# DELETEs it so the next bare tap doesn't replay. Public (no auth) on purpose:
# the booth page can be opened by anyone, locally or via myfaid.com (same
# backend through the tunnel), and arming alone does nothing without a real tap.
_iot_cart: Optional[dict] = None


@app.post("/api/iot/cart")
def set_iot_cart(body: IotCartRequest):
    global _iot_cart
    _iot_cart = body.model_dump()
    return {"status": "armed", "cart": _iot_cart}


@app.get("/api/iot/cart")
def get_iot_cart():
    return {"cart": _iot_cart}


@app.delete("/api/iot/cart")
def clear_iot_cart():
    global _iot_cart
    _iot_cart = None
    return {"status": "cleared"}


# Whitelist of sortable keys -> safe SQL expression. Never interpolate a raw
# client value into ORDER BY; only values in this map are allowed. Keys match
# the frontend's sort field keys 1:1.
_SORT_COLUMNS = {
    "transaction_id":  "transaction_id",
    "timestamp":       "timestamp",
    "customer_ref":    "customer_ref",
    "merchant_name":   "merchant_name",
    "amount_myr":      "amount_myr",
    "fraud_score":     "fraud_score",
    "predicted_label": "predicted_label",
    "reviewed_at":     "reviewed_at",
}


def _order_clause(sort, order):
    """Safe ORDER BY for the transaction list/export. Falls back to timestamp.
    NULLS LAST keeps unreviewed rows at the bottom; transaction_id is a stable
    tiebreaker so equal values don't reshuffle between requests."""
    col = _SORT_COLUMNS.get((sort or "").strip(), "timestamp")
    direction = "ASC" if (order or "").lower() == "asc" else "DESC"
    return f"ORDER BY {col} {direction} NULLS LAST, transaction_id ASC"


def _txn_filter(decision, reviewed, search, risk, customers=None):
    """Shared WHERE builder for transaction list / export / meta — returns
    (conditions, params). Risk bands: Low <0.6, Med 0.6–0.8, High ≥0.8.
    `customers` is an optional list of customer_refs to restrict to."""
    conditions, params = ["1=1"], []
    if decision:
        conditions.append("predicted_label = %s")
        params.append(decision.upper())
    if customers:
        placeholders = ",".join(["%s"] * len(customers))
        conditions.append(f"customer_ref IN ({placeholders})")
        params.extend(customers)
    if reviewed is True:
        conditions.append("reviewed_by IS NOT NULL AND reviewed_by != ''")
    elif reviewed is False:
        conditions.append("(reviewed_by IS NULL OR reviewed_by = '')")
    if search:
        conditions.append(
            "(transaction_id ILIKE %s OR merchant_name ILIKE %s OR customer_ref ILIKE %s)"
        )
        s = f"%{search}%"
        params.extend([s, s, s])
    if risk:
        r = risk.upper()
        if r == "LOW":
            conditions.append("fraud_score < 0.6")
        elif r == "MEDIUM":
            conditions.append("fraud_score >= 0.6 AND fraud_score < 0.8")
        elif r == "HIGH":
            conditions.append("fraud_score >= 0.8")
    return conditions, params


# Canonical export columns: key -> header. This list is also the output order.
_EXPORT_COLUMNS = [
    ("transaction_id",  "Transaction ID"),
    ("timestamp",       "Timestamp"),
    ("customer_ref",    "Customer Ref"),
    ("masked_card",     "Masked Card"),
    ("merchant",        "Merchant"),
    ("mcc",             "MCC"),
    ("mode",            "Mode"),
    ("location",        "Location"),
    ("amount",          "Amount (MYR)"),
    ("fraud_score",     "Fraud Score"),
    ("risk_level",      "Risk Level"),
    ("decision",        "Decision"),
    ("ml_prediction",   "ML Prediction"),
    ("rule_flag",       "Rule Flag"),
    ("rules_triggered", "Rules Triggered"),
    ("ground_truth",    "Ground Truth"),
    ("reviewed_by",     "Reviewed By"),
    ("reviewed_at",     "Reviewed At"),
]
_EXPORT_KEYS = [k for k, _ in _EXPORT_COLUMNS]

# Sensitive columns — only emitted for an authorized, consented PII export.
_PII_COLUMNS = [
    ("cardholder_name", "Cardholder Name"),
    ("ic_number",       "IC Number"),
    ("card_number",     "Card Number"),
    ("card_expiration", "Card Expiration"),
]
_PII_KEYS = [k for k, _ in _PII_COLUMNS]
_PII_ROLES = {"analyst", "admin"}   # roles allowed to export decrypted PII


def _require_pii_role(authorization: Optional[str]):
    """Validate the JWT and require an allowed role for a PII export."""
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Authentication required for PII export")
    try:
        payload = jwt.decode(authorization[7:], JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except JWTError:
        raise HTTPException(status_code=401, detail="Invalid or expired token")
    if payload.get("role") not in _PII_ROLES:
        raise HTTPException(status_code=403, detail="Your role is not permitted to export PII")
    return payload


@app.get("/api/export/transactions")
def export_transactions(
    decision: Optional[str] = Query(None),
    reviewed: Optional[bool] = Query(None),
    search: Optional[str] = Query(None),
    risk: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    columns: Optional[str] = Query(None),
    customers: Optional[str] = Query(None),
    sort: Optional[str] = Query(None),
    order: Optional[str] = Query(None),
    pii: bool = Query(False),
    format: str = Query("csv"),
    authorization: Optional[str] = Header(None),
):
    """Export transactions matching the same filters as GET /api/transactions
    (mirrors the table, all matching rows). `columns` (comma-separated keys),
    `date_from`/`date_to` (YYYY-MM-DD), `format` csv|xlsx. Privacy-safe by
    default; `pii=true` adds decrypted IC/card/cardholder and requires an
    authenticated user with an allowed role."""
    fmt = format.lower()
    if fmt not in ("csv", "xlsx"):
        raise HTTPException(status_code=400, detail="format must be 'csv' or 'xlsx'")

    if pii:
        _require_pii_role(authorization)

    # Columns available depend on whether PII was authorized
    available = _EXPORT_COLUMNS + (_PII_COLUMNS if pii else [])
    available_keys = [k for k, _ in available]
    headers_map = dict(available)

    if columns:
        requested = {c.strip() for c in columns.split(",") if c.strip()}
        selected = [k for k in available_keys if k in requested]
    else:
        selected = list(_EXPORT_KEYS)   # default = non-PII columns only
    if not selected:
        raise HTTPException(status_code=400, detail="No valid columns selected.")

    # Filters (shared with /api/transactions) + optional date range
    cust_list = [c.strip() for c in customers.split(",") if c.strip()] if customers else None
    conditions, params = _txn_filter(decision, reviewed, search, risk, cust_list)
    if date_from:
        conditions.append("timestamp::date >= %s")
        params.append(date_from)
    if date_to:
        conditions.append("timestamp::date <= %s")
        params.append(date_to)
    where = " AND ".join(conditions)

    base_cols = ("transaction_id, timestamp, customer_ref, masked_card_number, "
                 "merchant_name, mcc, mode, location, amount_myr, fraud_score, "
                 "predicted_label, ml_prediction, rule_flag, risk_reasons, "
                 "ground_truth_label, reviewed_by, reviewed_at")
    pii_select = (", cardholder_name, card_expiration_date, ic_number_enc, card_number_enc"
                  if pii else "")

    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute(
                f"SELECT {base_cols}{pii_select} FROM private_transactions "
                f"WHERE {where} {_order_clause(sort, order)}",
                params,
            )
            rows = cur.fetchall()
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    cipher = None
    if pii and ENCRYPTION_KEY:
        from cryptography.fernet import Fernet
        cipher = Fernet(ENCRYPTION_KEY)

    def _decrypt(value):
        if not value or not cipher:
            return ""
        try:
            return cipher.decrypt(value.encode()).decode()
        except Exception:
            return ""

    def _risk_level(score):
        if score is None:
            return ""
        if score >= 0.8:
            return "High"
        if score >= 0.6:
            return "Medium"
        return "Low"

    def _values(row):
        try:
            reasons = json.loads(row[13]) if row[13] else []
        except Exception:
            reasons = []
        reasons = [x for x in reasons if isinstance(x, str) and not x.startswith("ML score")]
        score = float(row[9]) if row[9] is not None else None
        v = {
            "transaction_id":  row[0],
            "timestamp":       row[1].isoformat() if row[1] else "",
            "customer_ref":    row[2] or "",
            "masked_card":     row[3] or "",
            "merchant":        row[4] or "",
            "mcc":             row[5] or "",
            "mode":            row[6] or "",
            "location":        row[7] or "",
            "amount":          round(float(row[8]), 2) if row[8] is not None else "",
            "fraud_score":     round(score, 4) if score is not None else "",
            "risk_level":      _risk_level(score),
            "decision":        row[10] or "",
            "ml_prediction":   int(row[11]) if row[11] is not None else 0,
            "rule_flag":       int(row[12]) if row[12] is not None else 0,
            "rules_triggered": "; ".join(reasons),
            "ground_truth":    int(row[14]) if row[14] is not None else "",
            "reviewed_by":     row[15] or "",
            "reviewed_at":     row[16].isoformat() if row[16] else "",
        }
        if pii:
            v["cardholder_name"] = row[17] or ""
            v["card_expiration"] = row[18] or ""
            v["ic_number"]       = _decrypt(row[19])
            v["card_number"]     = _decrypt(row[20])
        return v

    header   = [headers_map[k] for k in selected]
    rows_out = [_values(r) for r in rows]
    ts  = datetime.now().strftime("%Y%m%d-%H%M%S")
    tag = "transactions-pii" if pii else "transactions"

    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Transactions"
        ws.append(header)
        for v in rows_out:
            ws.append([v[k] for k in selected])
        bio = io.BytesIO()
        wb.save(bio)
        return Response(
            content=bio.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{tag}-{ts}.xlsx"'},
        )

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    for v in rows_out:
        writer.writerow([v[k] for k in selected])
    return Response(
        content=buf.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{tag}-{ts}.csv"'},
    )


@app.get("/api/export/transactions/meta")
def export_meta(
    decision: Optional[str] = Query(None),
    reviewed: Optional[bool] = Query(None),
    search: Optional[str] = Query(None),
    risk: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    customers: Optional[str] = Query(None),
):
    """For the export modal: how many rows the export will contain (filters +
    date range) and the selectable date window (min/max date over the filters,
    ignoring the date range). Used to size the count + constrain the calendars."""
    cust_list = [c.strip() for c in customers.split(",") if c.strip()] if customers else None
    conditions, params = _txn_filter(decision, reviewed, search, risk, cust_list)
    filter_where = " AND ".join(conditions)

    # COUNT applies the date range; MIN/MAX define the selectable bounds (no dates).
    date_clause, date_params = "TRUE", []
    if date_from:
        date_clause += " AND timestamp::date >= %s"
        date_params.append(date_from)
    if date_to:
        date_clause += " AND timestamp::date <= %s"
        date_params.append(date_to)

    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT COUNT(*) FILTER (WHERE {date_clause}),
                       MIN(timestamp)::date,
                       MAX(timestamp)::date
                FROM private_transactions
                WHERE {filter_where}
                """,
                date_params + params,
            )
            count, min_d, max_d = cur.fetchone()
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    return {
        "count":    int(count or 0),
        "min_date": min_d.isoformat() if min_d else None,
        "max_date": max_d.isoformat() if max_d else None,
    }


@app.get("/api/transactions/customer/{customer_ref}")
def get_customer_transactions(customer_ref: str, limit: int = Query(20, le=50)):
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT transaction_id, customer_ref, timestamp, amount_myr,
                       merchant_name, mcc, mode, location, ic_hash,
                       masked_card_number, fraud_score, predicted_label,
                       ml_prediction, rule_flag, risk_reasons,
                       ground_truth_label, reviewed_by, reviewed_at, created_at
                FROM private_transactions
                WHERE customer_ref = %s
                ORDER BY timestamp DESC
                LIMIT %s
                """,
                (customer_ref, limit),
            )
            rows = cur.fetchall()
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    def _row_to_txn(row):
        try:
            risk_reasons = json.loads(row[14]) if row[14] else []
        except Exception:
            risk_reasons = []
        return {
            "transaction_id":     row[0],
            "customer_ref":       row[1] or "",
            "timestamp":          row[2].isoformat() if row[2] else "",
            "amount_myr":         float(row[3]) if row[3] is not None else 0.0,
            "merchant_name":      row[4] or "",
            "mcc":                row[5] or "",
            "mode":               row[6] or "",
            "location":           row[7] or "",
            "ic_hash":            row[8] or "",
            "masked_card_number": row[9] or "",
            "fraud_score":        float(row[10]) if row[10] is not None else 0.0,
            "predicted_label":    row[11] or "pending",
            "ml_prediction":      int(row[12]) if row[12] is not None else 0,
            "rule_flag":          int(row[13]) if row[13] is not None else 0,
            "risk_reasons":       risk_reasons,
            "ground_truth_label": int(row[15]) if row[15] is not None else 0,
            "reviewed_by":        row[16] or "",
            "reviewed_at":        row[17].isoformat() if row[17] else "",
            "created_at":         row[18].isoformat() if row[18] else "",
            "created_by":         "",
        }

    txns = [_row_to_txn(r) for r in rows]
    return {"customer_ref": customer_ref, "total": len(txns), "transactions": txns}


@app.get("/api/transactions/{transaction_id}")
def get_transaction(transaction_id: str):
    try:
        blockchain_data = fabric.get_transaction(transaction_id)
    except Exception as e:
        raise HTTPException(status_code=404, detail=f"Transaction not found: {e}")

    try:
        from cryptography.fernet import Fernet, InvalidToken
        cipher = Fernet(ENCRYPTION_KEY) if ENCRYPTION_KEY else None

        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT pt.cardholder_name, pt.card_expiration_date,
                       pt.ic_number_enc, pt.card_number_enc,
                       pt.reviewed_by, pt.reviewed_at, pt.notes,
                       fc.frozen_at, fc.trigger_reasons
                FROM private_transactions pt
                LEFT JOIN frozen_cards fc ON fc.card_hash = pt.card_hash
                WHERE pt.transaction_id = %s
            """, (transaction_id,))
            row = cur.fetchone()
        conn.close()
        if row:
            blockchain_data["cardholder_name"]     = row[0]
            blockchain_data["card_expiration_date"] = row[1]

            def _decrypt(value: str) -> str | None:
                if not value or not cipher:
                    return None
                try:
                    return cipher.decrypt(value.encode()).decode()
                except (InvalidToken, Exception):
                    return None

            blockchain_data["ic_number"]           = _decrypt(row[2])
            blockchain_data["card_number"]         = _decrypt(row[3])
            blockchain_data["private_reviewed_by"] = row[4]
            blockchain_data["private_reviewed_at"] = str(row[5]) if row[5] else ""
            blockchain_data["notes"]               = row[6]

            frozen_at = row[7]
            blockchain_data["card_frozen"]    = frozen_at is not None
            blockchain_data["card_frozen_at"] = frozen_at.isoformat() if frozen_at else ""
            try:
                blockchain_data["card_frozen_reasons"] = json.loads(row[8]) if row[8] else []
            except Exception:
                blockchain_data["card_frozen_reasons"] = []
    except Exception as e:
        blockchain_data["private_data_error"] = str(e)

    return blockchain_data


@app.get("/api/transactions/{transaction_id}/history")
def get_transaction_history(transaction_id: str):
    try:
        return fabric.get_history(transaction_id)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Fabric error: {e}")


@app.post("/api/transactions/{transaction_id}/review")
def review_transaction(transaction_id: str, body: ReviewRequest):
    if body.ground_truth_label not in [0, 1]:
        raise HTTPException(status_code=400, detail="ground_truth_label must be 0 or 1")
    if not body.reviewer_id.strip():
        raise HTTPException(status_code=400, detail="reviewer_id is required")

    try:
        fabric.update_ground_truth(
            transaction_id=transaction_id,
            ground_truth_label=body.ground_truth_label,
            reviewer_id=body.reviewer_id,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Fabric update failed: {e}")

    try:
        conn = get_db()
        private_store = PrivateRecordStore(conn, ENCRYPTION_KEY)
        private_store.update_ground_truth(
            transaction_id=transaction_id,
            ground_truth_label=body.ground_truth_label,
            reviewer_id=body.reviewer_id,
            notes=body.notes or "",
        )
        # A review can clear a false-positive fraud → recompute the card's
        # freeze (unfreezes when no standing fraud remains on the card).
        with conn.cursor() as cur:
            cur.execute(
                "SELECT card_hash, merchant_name, amount_myr "
                "FROM private_transactions WHERE transaction_id = %s",
                (transaction_id,),
            )
            r = cur.fetchone()
        unfroze = False
        if r and r[0]:
            unfroze = private_store.reconcile_card_freeze(r[0])
        private_store.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Postgres update failed: {e}")

    # Notify all SSE clients. Clearing a false-positive that unfreezes the card
    # gets its own event so the UI can announce it (mirrors the freeze toast);
    # otherwise it's a plain review (both trigger a dashboard refresh).
    if unfroze:
        _notify_clients("unfreeze", {
            "transaction_id": transaction_id,
            "merchant_name":  r[1] or "",
            "amount_myr":     float(r[2]) if r[2] is not None else 0.0,
        })
    else:
        _notify_clients("review", {"transaction_id": transaction_id})

    return {"status": "SUCCESS", "transaction_id": transaction_id}


# ── Customers (KYC) ──────────────────────────────────────────

@app.get("/api/customers/search")
def search_customers(q: str = Query("", min_length=0), limit: int = Query(10, le=50)):
    """Typeahead for the export/table customer filter. Matches customer_ref
    (case-insensitive prefix/substring). Returns ref + name only — no PII.
    Registered before /{customer_ref} so the literal path wins routing."""
    term = q.strip()
    if not term:
        return []
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT customer_ref, name
                FROM kyc_profiles
                WHERE customer_ref ILIKE %s
                ORDER BY customer_ref
                LIMIT %s
                """,
                (f"%{term}%", limit),
            )
            rows = cur.fetchall()
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")
    return [{"customer_ref": r[0], "name": r[1]} for r in rows]


@app.get("/api/customers/{customer_ref}")
def get_customer_kyc(customer_ref: str):
    conn = get_db()
    try:
        store   = KYCStore(conn, ENCRYPTION_KEY)
        profile = store.get_by_customer_ref(customer_ref)
        if not profile:
            raise HTTPException(status_code=404, detail=f"Customer {customer_ref} not found")
        return profile
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")
    finally:
        conn.close()


# ── Stats ─────────────────────────────────────────────────────

@app.get("/api/stats")
def get_stats():
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT
                    COUNT(*)                                                            AS total,
                    COUNT(*) FILTER (WHERE predicted_label = 'FRAUD')                  AS fraud_count,
                    COUNT(*) FILTER (WHERE predicted_label = 'LEGIT')                  AS legit_count,
                    COUNT(*) FILTER (WHERE predicted_label = 'FRAUD'
                                     AND reviewed_by IS NULL)                          AS pending_review,
                    COUNT(*) FILTER (WHERE predicted_label = 'FRAUD'
                                     AND reviewed_by IS NOT NULL)                      AS reviewed,
                    COALESCE(AVG(fraud_score), 0)                                      AS avg_fraud_score,
                    COALESCE(SUM(amount_myr) FILTER (WHERE predicted_label = 'FRAUD'), 0) AS balance_at_risk
                FROM private_transactions
            """)
            total, fraud_count, legit_count, pending, reviewed, avg_score, balance_at_risk = cur.fetchone()
            cur.execute("SELECT COUNT(*) FROM frozen_cards")
            compromised_cards = cur.fetchone()[0]
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    return {
        "total":             int(total),
        "fraud_count":       int(fraud_count),
        "legit_count":       int(legit_count),
        "pending_review":    int(pending),
        "reviewed":          int(reviewed),
        "avg_fraud_score":   round(float(avg_score), 4),
        "balance_at_risk":   float(balance_at_risk),
        "compromised_cards": int(compromised_cards),
    }


# ── Frozen cards ──────────────────────────────────────────────

def _json_list(value):
    try:
        return json.loads(value) if value else []
    except Exception:
        return []


@app.get("/api/frozen-cards")
def list_frozen_cards():
    """List every auto-frozen card with per-card fraud count + amount at risk,
    plus a summary (total, frozen in the latest data month, total at risk)."""
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT fc.card_hash, fc.customer_ref, fc.cardholder_name, fc.card_last4,
                       fc.frozen_at, fc.trigger_txn_id, fc.trigger_reasons,
                       COUNT(pt.transaction_id) FILTER (WHERE pt.predicted_label = 'FRAUD') AS fraud_count,
                       COALESCE(SUM(pt.amount_myr) FILTER (WHERE pt.predicted_label = 'FRAUD'), 0) AS amount_at_risk,
                       COUNT(pt.transaction_id) AS total_txns
                FROM frozen_cards fc
                LEFT JOIN private_transactions pt ON pt.card_hash = fc.card_hash
                GROUP BY fc.card_hash, fc.customer_ref, fc.cardholder_name, fc.card_last4,
                         fc.frozen_at, fc.trigger_txn_id, fc.trigger_reasons
                ORDER BY fc.frozen_at DESC
            """)
            rows = cur.fetchall()
            cur.execute("""
                SELECT COUNT(*),
                       COUNT(*) FILTER (
                         WHERE DATE_TRUNC('month', frozen_at) =
                               (SELECT DATE_TRUNC('month', MAX(frozen_at)) FROM frozen_cards)
                       )
                FROM frozen_cards
            """)
            total, this_month = cur.fetchone()
            cur.execute(
                "SELECT COALESCE(SUM(amount_myr) FILTER (WHERE predicted_label = 'FRAUD'), 0) "
                "FROM private_transactions"
            )
            total_at_risk = cur.fetchone()[0]
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    cards = [{
        "card_hash":       r[0],
        "customer_ref":    r[1] or "",
        "cardholder_name": r[2] or "",
        "card_last4":      r[3] or "",
        "frozen_at":       r[4].isoformat() if r[4] else "",
        "trigger_txn_id":  r[5] or "",
        "trigger_reasons": _json_list(r[6]),
        "fraud_count":     int(r[7]),
        "amount_at_risk":  float(r[8]),
        "total_txns":      int(r[9]),
    } for r in rows]

    return {
        "summary": {
            "total":         int(total or 0),
            "this_month":    int(this_month or 0),
            "total_at_risk": float(total_at_risk or 0),
        },
        "cards": cards,
    }


@app.get("/api/frozen-cards/{card_hash}")
def get_frozen_card(card_hash: str):
    """One frozen card + all its transactions oldest-first, each tagged
    is_post_freeze (timestamp after the freeze) so the UI can draw a
    'frozen here' divider. 404 if the card isn't frozen."""
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT card_hash, customer_ref, cardholder_name, card_last4,
                       frozen_at, trigger_txn_id, trigger_reasons
                FROM frozen_cards WHERE card_hash = %s
            """, (card_hash,))
            card = cur.fetchone()
            if not card:
                raise HTTPException(status_code=404, detail="Card is not frozen")
            cur.execute("""
                SELECT transaction_id, timestamp, amount_myr, merchant_name,
                       mcc, mode, location, fraud_score, predicted_label,
                       risk_reasons, reviewed_by
                FROM private_transactions
                WHERE card_hash = %s
                ORDER BY timestamp ASC
            """, (card_hash,))
            txn_rows = cur.fetchall()
        conn.close()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    frozen_at = card[4]
    txns = []
    for t in txn_rows:
        ts = t[1]
        txns.append({
            "transaction_id":  t[0],
            "timestamp":       ts.isoformat() if ts else "",
            "amount_myr":      float(t[2]) if t[2] is not None else 0.0,
            "merchant_name":   t[3] or "",
            "mcc":             t[4] or "",
            "mode":            t[5] or "",
            "location":        t[6] or "",
            "fraud_score":     float(t[7]) if t[7] is not None else 0.0,
            "predicted_label": t[8] or "pending",
            "risk_reasons":    _json_list(t[9]),
            "reviewed_by":     t[10] or "",
            "is_post_freeze":  bool(frozen_at and ts and ts > frozen_at),
        })

    return {
        "card_hash":       card[0],
        "customer_ref":    card[1] or "",
        "cardholder_name": card[2] or "",
        "card_last4":      card[3] or "",
        "frozen_at":       frozen_at.isoformat() if frozen_at else "",
        "trigger_txn_id":  card[5] or "",
        "trigger_reasons": _json_list(card[6]),
        "transactions":    txns,
    }


# ── Charts ────────────────────────────────────────────────────

@app.get("/api/charts/fraud-trend")
def fraud_trend():
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT TO_CHAR(DATE_TRUNC('month', timestamp), 'YYYY-MM') AS month,
                       COUNT(*) FILTER (WHERE predicted_label = 'FRAUD')                     AS fraud,
                       COUNT(*) FILTER (WHERE predicted_label = 'LEGIT')                     AS legit,
                       COALESCE(SUM(amount_myr) FILTER (WHERE predicted_label = 'FRAUD'), 0) AS amount_at_risk
                FROM private_transactions
                WHERE predicted_label IN ('FRAUD', 'LEGIT')
                GROUP BY DATE_TRUNC('month', timestamp)
                ORDER BY 1
            """)
            rows = cur.fetchall()
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    return [
        {"date": m, "fraud": int(f), "legit": int(l), "amount_at_risk": float(a)}
        for m, f, l, a in rows
    ]


@app.get("/api/charts/score-distribution")
def score_distribution():
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT
                    COUNT(*) FILTER (WHERE fraud_score < 0.6)                        AS low,
                    COUNT(*) FILTER (WHERE fraud_score >= 0.6 AND fraud_score < 0.8) AS medium,
                    COUNT(*) FILTER (WHERE fraud_score >= 0.8)                       AS high
                FROM private_transactions
                WHERE predicted_label = 'FRAUD' AND fraud_score IS NOT NULL
            """)
            low, medium, high = cur.fetchone()
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    return [
        {"range": "Low",    "count": int(low)},
        {"range": "Medium", "count": int(medium)},
        {"range": "High",   "count": int(high)},
    ]


# ── Detection triggers (rule firing counts + ML feature importance) ──

# Parent names of one-hot encoded categorical features. preprocess_data()
# get_dummies() produces "<Parent>_<value>" columns; we sum a parent's dummies
# back together so the chart shows ~12 meaningful signals, not fragmented dummies.
_CATEGORICAL_PARENTS = [
    "Date of Birth", "Employment Status", "Job Title", "Marital Status",
    "Device Information", "IP Address", "Location", "City", "State",
    "Country", "Nationality", "Gender", "Mode", "Card Type",
]

# Friendly labels + group for numeric (non-dummy) features. Anything not here
# that isn't a categorical parent falls back to its raw name / "Other".
_NUMERIC_FEATURES = {
    "hours_since_last_txn":   ("Hours since last txn",        "Velocity"),
    "amount_zscore_cust":     ("Amount vs customer norm (z)", "Velocity"),
    "amount_vs_cust_avg":     ("Amount vs customer average",  "Velocity"),
    "cust_avg_amount":        ("Customer average amount",     "Velocity"),
    "cust_std_amount":        ("Customer amount spread",      "Velocity"),
    "txn_count_history":      ("Customer history depth",      "Velocity"),
    "new_mcc_flag":           ("New merchant category",       "Velocity"),
    "location_changed":       ("Location changed",            "Velocity"),
    "log_amount":             ("Amount (log)",                "Engineered"),
    "amount_is_round":        ("Round-number amount",         "Engineered"),
    "amount_to_income_ratio": ("Amount-to-income ratio",      "Engineered"),
    "has_income_data":        ("Has income data",             "Engineered"),
    "income_min":             ("Income (min)",                "Engineered"),
    "income_max":             ("Income (max)",                "Engineered"),
    "income_mid":             ("Income (midpoint)",           "Engineered"),
    "Amount (MYR)":           ("Transaction amount",          "Transaction"),
    "txn_hour":               ("Hour of day",                 "Transaction"),
    "txn_dayofweek":          ("Day of week",                 "Transaction"),
    "Merchant Category Code (MCC)": ("Merchant category (MCC)", "Transaction"),
}

_ml_importance_cache: Optional[list] = None


def _parent_of(feature: str) -> str:
    for parent in _CATEGORICAL_PARENTS:
        if feature == parent or feature.startswith(parent + "_"):
            return parent
    return feature


def _compute_ml_importances(top_n: int = 12) -> list:
    global _ml_importance_cache
    if _ml_importance_cache is not None:
        return _ml_importance_cache

    model = joblib.load("models/fraud_model.pkl")
    cols  = joblib.load("models/feature_columns.pkl")

    agg: dict[str, float] = defaultdict(float)
    for col, imp in zip(cols, model.feature_importances_):
        agg[_parent_of(col)] += float(imp)

    ranked = sorted(agg.items(), key=lambda kv: kv[1], reverse=True)[:top_n]

    result = []
    for raw, importance in ranked:
        if raw in _NUMERIC_FEATURES:
            label, group = _NUMERIC_FEATURES[raw]
        elif raw in _CATEGORICAL_PARENTS:
            label, group = raw, "Profile"
        else:
            label, group = raw, "Other"
        result.append({
            "label":      label,
            "raw":        raw,
            "group":      group,
            "importance": round(importance, 4),
        })

    _ml_importance_cache = result
    return result


@app.get("/api/triggers/stats")
def trigger_stats():
    """Live data for the Detection Triggers page: how often each rule has
    fired across scored transactions + aggregated ML feature importances."""
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*), COUNT(*) FILTER (WHERE predicted_label = 'FRAUD') FROM private_transactions")
            total, total_fraud = cur.fetchone()

            cur.execute("SELECT risk_reasons FROM private_transactions WHERE risk_reasons IS NOT NULL")
            rows = cur.fetchall()
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    rule_counts: dict[str, int] = defaultdict(int)
    for (rr,) in rows:
        reasons = json.loads(rr) if isinstance(rr, str) else (rr or [])
        for reason in reasons:
            # skip the "ML score = 0.xxxx" entries written for non-rule rows
            if isinstance(reason, str) and not reason.startswith("ML score"):
                rule_counts[reason] += 1

    try:
        importances = _compute_ml_importances()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Model load error: {e}")

    return {
        "total_transactions": int(total),
        "total_fraud":        int(total_fraud),
        "rule_counts":        dict(rule_counts),
        "ml": {
            "model":       "Random Forest",
            "n_features":  93,
            "importances": importances,
        },
    }


def _metrics_from_counts(tp: int, fp: int, tn: int, fn: int) -> dict:
    """Precision/recall/F1/FPR/accuracy from a confusion matrix (no AUC —
    that needs probabilities, which the live DB path does not store)."""
    precision = tp / (tp + fp) if (tp + fp) else 0.0
    recall    = tp / (tp + fn) if (tp + fn) else 0.0
    f1        = (2 * precision * recall / (precision + recall)
                 if (precision + recall) else 0.0)
    fpr       = fp / (fp + tn) if (fp + tn) else 0.0
    total     = tp + fp + tn + fn
    accuracy  = (tp + tn) / total if total else 0.0
    return {
        "accuracy":  round(accuracy, 4),
        "precision": round(precision, 4),
        "recall":    round(recall, 4),
        "f1":        round(f1, 4),
        "fpr":       round(fpr, 4),
    }


@app.get("/api/model/performance")
def model_performance():
    """Detection-model performance for the Triggers page modal.

    `test` — honest held-out evaluation of the *deployed* model, generated
    offline by `python -m src.eval_model` (20% stratified split the model never
    saw during training); read from models/model_metrics.json.

    `live` — the same model's classification on every scored transaction that
    has a ground-truth label, computed from the DB right now (ml_prediction vs
    ground_truth_label). No live AUC: the ML probability is not persisted.
    """
    # ── Held-out test metrics (from the offline eval artifact) ──
    test = None
    try:
        with open("models/model_metrics.json") as fh:
            test = json.load(fh)
    except FileNotFoundError:
        test = None
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Could not read metrics file: {e}")

    # ── Live metrics from the DB (ML prediction vs ground truth) ──
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT
                    COUNT(*) FILTER (WHERE ml_prediction = 1 AND ground_truth_label = 1),
                    COUNT(*) FILTER (WHERE ml_prediction = 1 AND ground_truth_label = 0),
                    COUNT(*) FILTER (WHERE ml_prediction = 0 AND ground_truth_label = 0),
                    COUNT(*) FILTER (WHERE ml_prediction = 0 AND ground_truth_label = 1)
                FROM private_transactions
                WHERE ground_truth_label IS NOT NULL AND ml_prediction IS NOT NULL
            """)
            tp, fp, tn, fn = (int(v) for v in cur.fetchone())
        conn.close()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Database error: {e}")

    live = {
        "n_scored":   tp + fp + tn + fn,
        "positives":  tp + fn,
        "negatives":  tn + fp,
        "confusion_matrix": {"tp": tp, "fp": fp, "tn": tn, "fn": fn},
        **_metrics_from_counts(tp, fp, tn, fn),
    }

    return {"test": test, "live": live}


# ── Audit: chain vs DB integrity ─────────────────────────────

@app.get("/api/audit/integrity-check")
def integrity_check():
    """
    Compares every record on the Fabric blockchain against the private
    PostgreSQL store and flags any count or field-level discrepancies.
    Fields checked: fraud_score, predicted_label, amount_myr,
                    ml_prediction, rule_flag.
    """
    try:
        fabric_result = fabric.get_all_transactions()
        fabric_txns = fabric_result if isinstance(fabric_result, list) else fabric_result.get("transactions", [])
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"Fabric unavailable: {e}")

    fabric_map: dict = {
        t["transaction_id"]: t
        for t in fabric_txns
        if t.get("transaction_id")
    }

    conn = get_db()
    db_map: dict = {}
    try:
        with conn.cursor() as cur:
            cur.execute("""
                SELECT transaction_id, fraud_score, predicted_label,
                       amount_myr, ml_prediction, rule_flag
                FROM private_transactions
            """)
            for row in cur.fetchall():
                db_map[row[0]] = {
                    "fraud_score":     float(row[1]) if row[1] is not None else None,
                    "predicted_label": row[2],
                    "amount_myr":      float(row[3]) if row[3] is not None else None,
                    "ml_prediction":   int(row[4])   if row[4] is not None else None,
                    "rule_flag":       int(row[5])   if row[5] is not None else None,
                }
    finally:
        conn.close()

    fabric_ids  = set(fabric_map.keys())
    db_ids      = set(db_map.keys())
    common      = fabric_ids & db_ids
    only_chain  = fabric_ids - db_ids
    only_db     = db_ids - fabric_ids

    FLOAT_FIELDS   = {"fraud_score", "amount_myr"}
    COMPARE_FIELDS = ["fraud_score", "predicted_label", "amount_myr", "ml_prediction", "rule_flag"]

    mismatch_details = []
    matched = 0
    for tid in common:
        f_rec = fabric_map[tid]
        d_rec = db_map[tid]
        record_ok = True
        for field in COMPARE_FIELDS:
            f_val = f_rec.get(field)
            d_val = d_rec.get(field)
            if field in FLOAT_FIELDS:
                try:
                    match = abs(float(f_val) - float(d_val)) <= 0.001
                except (TypeError, ValueError):
                    match = (f_val == d_val)
            else:
                match = str(f_val).upper() == str(d_val).upper()
            if not match:
                mismatch_details.append({
                    "transaction_id": tid,
                    "field":  field,
                    "chain":  f_val,
                    "db":     d_val,
                })
                record_ok = False
                break
        if record_ok:
            matched += 1

    tampered = bool(mismatch_details) or bool(only_db) or bool(only_chain)

    return {
        "status":          "tampered" if tampered else "verified",
        "chain_total":     len(fabric_ids),
        "db_total":        len(db_ids),
        "checked":         len(common),
        "matched":         matched,
        "mismatches":      len(mismatch_details),
        "only_in_chain":   len(only_chain),
        "only_in_db":      len(only_db),
        "mismatch_sample": mismatch_details[:3],
        "checked_at":      datetime.now(timezone.utc).isoformat(),
    }


# ── Blockchain network view ──────────────────────────────────
#
# Fabric nodes expose an operations service (plain HTTP in this test-network):
# /healthz for liveness and /metrics (Prometheus) for ledger block height.
# Topology is fixed by the test-network compose; ports verified against the
# running containers (orderer 9443, peer0.org1 9444, peer0.org2 9445).

_OPS_HOST = os.getenv("FABRIC_OPS_HOST", "localhost")
_BLOCKCHAIN_NODES = [
    {"id": "orderer",    "name": "Orderer",      "role": "orderer", "org": "OrdererOrg", "msp": "OrdererMSP", "endpoint": "orderer.example.com:7050",   "ops_port": 9443},
    {"id": "peer0-org1", "name": "Peer0 · Org1", "role": "peer",    "org": "Org1",       "msp": "Org1MSP",    "endpoint": "peer0.org1.example.com:7051", "ops_port": 9444},
    {"id": "peer0-org2", "name": "Peer0 · Org2", "role": "peer",    "org": "Org2",       "msp": "Org2MSP",    "endpoint": "peer0.org2.example.com:9051", "ops_port": 9445},
]
_BLOCK_HEIGHT_RE = re.compile(r'^ledger_blockchain_height\{[^}]*\}\s+([0-9.e+]+)', re.MULTILINE)


def _scrape_block_height(ops_port: int) -> Optional[int]:
    """Pull ledger_blockchain_height for the channel from a node's /metrics."""
    try:
        r = requests.get(f"http://{_OPS_HOST}:{ops_port}/metrics", timeout=2)
        if r.status_code != 200:
            return None
        vals = [int(float(v)) for v in _BLOCK_HEIGHT_RE.findall(r.text)]
        return max(vals) if vals else None
    except Exception:
        return None


@app.get("/api/blockchain/nodes")
def blockchain_nodes():
    """Live network health: per-node liveness (/healthz) + block height (/metrics),
    plus gateway reachability. Powers the Blockchain page network view."""
    nodes = []
    heights = []
    for spec in _BLOCKCHAIN_NODES:
        status, latency_ms = "down", None
        try:
            t0 = time.perf_counter()
            r = requests.get(f"http://{_OPS_HOST}:{spec['ops_port']}/healthz", timeout=2)
            latency_ms = round((time.perf_counter() - t0) * 1000, 1)
            status = "up" if r.status_code == 200 else "unhealthy"
        except Exception:
            status = "down"
        height = _scrape_block_height(spec["ops_port"])
        if height is not None:
            heights.append(height)
        nodes.append({**spec, "status": status, "latency_ms": latency_ms, "block_height": height})

    gateway_ok = False
    try:
        requests.get(f"{FABRIC_GATEWAY_URL}/all", timeout=3).raise_for_status()
        gateway_ok = True
    except Exception:
        gateway_ok = False

    up = sum(1 for n in nodes if n["status"] == "up")
    return {
        "nodes":         nodes,
        "nodes_up":      up,
        "nodes_total":   len(nodes),
        "all_healthy":   up == len(nodes) and gateway_ok,
        "channel":       "mychannel",
        "chaincode":     os.getenv("CHAINCODE_NAME", "fraud"),
        "block_height":  max(heights) if heights else None,
        "gateway_ok":    gateway_ok,
        "checked_at":    datetime.now(timezone.utc).isoformat(),
    }


@app.get("/api/blockchain/chain")
def blockchain_chain(blocks: int = Query(8, ge=1, le=50)):
    """Real ledger hash chain from qscc: tip info + the last N blocks
    (number, data hash, previous-block hash) — proof the chain is append-only."""
    try:
        info = fabric.get_chain_info()
        recent = fabric.get_blocks(count=blocks)
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"Fabric unavailable: {e}")
    return {
        "channel":             info.get("channel", "mychannel"),
        "height":              info.get("height"),
        "current_block_hash":  info.get("current_block_hash"),
        "previous_block_hash": info.get("previous_block_hash"),
        "blocks":              recent.get("blocks", []),
        "checked_at":          datetime.now(timezone.utc).isoformat(),
    }


# ── Internal endpoints: local-only guard ─────────────────────
# /internal/* must never be reachable through the public tunnel. A localhost
# client.host is NOT enough — cloudflared forwards from localhost, so every
# tunneled request looks local (127.0.0.1). Cloudflare/any proxy adds a
# forwarding header the caller can't remove, so reject whenever one is present.
_FORWARD_HEADERS = ("x-forwarded-for", "cf-connecting-ip", "x-real-ip", "forwarded")


def require_local(request: Request):
    client_host = request.client.host if request.client else ""
    if client_host not in ("127.0.0.1", "::1", "localhost") or \
       any(h in request.headers for h in _FORWARD_HEADERS):
        raise HTTPException(status_code=403, detail="Internal endpoint — local access only")


# ── Internal reset (demo only) ───────────────────────────────

@app.post("/internal/reset-demo")
def reset_demo(_local: None = Depends(require_local)):
    conn = get_db()
    with conn.cursor() as cur:
        # frozen_cards is derived from private_transactions — wipe both so a
        # nuclear reset doesn't leave orphaned freezes pointing at gone rows.
        cur.execute("TRUNCATE TABLE private_transactions, frozen_cards RESTART IDENTITY")
    conn.commit()
    conn.close()
    _notify_clients("reset", {})
    return {"status": "cleared"}


# ── Run / reset the demo deck ────────────────────────────────
# These manage just the curated demo data (is_demo=true rows), leaving the
# historical backfill untouched.

import subprocess
import sys

_demo_proc: subprocess.Popen | None = None


def _demo_running() -> bool:
    global _demo_proc
    return _demo_proc is not None and _demo_proc.poll() is None


@app.post("/internal/run-demo")
def run_demo(delay: float = 2.0, _local: None = Depends(require_local)):
    """
    Spawn the demo deck producer as a background subprocess. Plays through
    the deck once and exits. If a demo is already running, returns 409.
    """
    global _demo_proc
    if _demo_running():
        raise HTTPException(status_code=409, detail="Demo already running")

    backend_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    cmd = [sys.executable, "-m", "src.kafka_producer", "--demo",
           "--delay", str(delay)]

    _demo_proc = subprocess.Popen(
        cmd,
        cwd=backend_dir,
        env={**os.environ},
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return {"status": "started", "pid": _demo_proc.pid, "delay": delay}


@app.post("/internal/stop-demo")
def stop_demo(_local: None = Depends(require_local)):
    """Stops the running demo producer (if any). Idempotent."""
    global _demo_proc
    if not _demo_running():
        return {"status": "not_running"}
    _demo_proc.terminate()
    try:
        _demo_proc.wait(timeout=3)
    except subprocess.TimeoutExpired:
        _demo_proc.kill()
    return {"status": "stopped"}


@app.get("/internal/demo-status")
def demo_status(_local: None = Depends(require_local)):
    """Returns whether the demo producer is currently running."""
    return {"running": _demo_running(),
            "pid": _demo_proc.pid if _demo_running() else None}


@app.post("/internal/reset-demo-data")
def reset_demo_data(_local: None = Depends(require_local)):
    """
    Delete only the demo rows (is_demo=true). Historical data is untouched.
    Also stops the demo if it's currently running.
    """
    global _demo_proc
    if _demo_running():
        _demo_proc.terminate()
        try:
            _demo_proc.wait(timeout=3)
        except subprocess.TimeoutExpired:
            _demo_proc.kill()

    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM private_transactions WHERE is_demo = TRUE")
            deleted = cur.rowcount
            # Demo timestamps always fall after all historical data, so a demo
            # fraud is never a card's *earliest* fraud — any frozen_cards row
            # triggered by a (now-deleted) demo txn belongs to a card with no
            # remaining fraud. Drop those orphans; historical-triggered freezes
            # keep their still-present trigger row and are untouched.
            cur.execute("""
                DELETE FROM frozen_cards
                WHERE trigger_txn_id NOT IN (
                    SELECT transaction_id FROM private_transactions
                )
            """)
            unfrozen = cur.rowcount
        conn.commit()
    finally:
        conn.close()

    _notify_clients("reset", {})
    return {"status": "cleared", "deleted": deleted, "unfrozen": unfrozen}


# ── Internal notify (called by Kafka consumer) ────────────────

@app.post("/internal/notify")
async def internal_notify(request: Request, _local: None = Depends(require_local)):
    """
    Called by kafka_fraud_consumer after saving a new transaction.
    Triggers an immediate SSE push to all connected dashboard clients.
    Restricted to local access only (see require_local).
    """
    body = await request.json()
    _notify_clients("new_transaction", body)
    return {"notified": len(_sse_clients)}


# ── SSE Stream ────────────────────────────────────────────────

@app.get("/api/stream")
async def stream_dashboard(request: Request):
    """
    Event-driven SSE endpoint.

    - Sends a full snapshot immediately on connect
    - Pushes a fresh snapshot when:
        (a) Kafka consumer saves a new transaction  → POST /internal/notify
        (b) A reviewer submits a decision           → POST /api/transactions/{id}/review
    - Sends a keepalive comment every 30s so the
      connection isn't dropped by proxies/browsers
    - No fixed polling interval — only fires when data changes
    """
    queue: asyncio.Queue = asyncio.Queue(maxsize=20)
    _sse_clients.add(queue)

    async def event_generator():
        try:
            # Send initial snapshot on connect
            try:
                snapshot = await asyncio.get_event_loop().run_in_executor(
                    None, _build_snapshot
                )
                yield f"event: dashboard\ndata: {json.dumps(snapshot)}\n\n"
            except Exception as e:
                # Keep the stream alive; client will retry on next event/keepalive
                print(f"[SSE] initial snapshot error: {e}", flush=True)
                yield ": keepalive\n\n"

            # Wait for events, push snapshot on each one
            while True:
                if await request.is_disconnected():
                    break

                try:
                    # Block until an event arrives or 30s keepalive timeout
                    event_data = await asyncio.wait_for(queue.get(), timeout=30)

                    # Reset event: notify the client then close this connection.
                    # The browser's EventSource will reconnect automatically and
                    # get a fresh empty snapshot — no manual reconnect needed.
                    if isinstance(event_data, dict) and event_data.get("type") == "reset":
                        yield "event: reset\ndata: {}\n\n"
                        return

                    # Pass-through events the client reacts to directly (e.g. an
                    # unfreeze toast). Collect from this event and the drain below.
                    passthrough = []
                    if isinstance(event_data, dict) and event_data.get("type") == "unfreeze":
                        passthrough.append(event_data)

                    # Debounce: drain any events that stacked up while
                    # we were building the last snapshot
                    while not queue.empty():
                        ev = queue.get_nowait()
                        if isinstance(ev, dict) and ev.get("type") == "unfreeze":
                            passthrough.append(ev)

                    for ev in passthrough:
                        yield f"event: unfreeze\ndata: {json.dumps(ev)}\n\n"

                    try:
                        snapshot = await asyncio.get_event_loop().run_in_executor(
                            None, _build_snapshot
                        )
                        yield f"event: dashboard\ndata: {json.dumps(snapshot)}\n\n"
                    except Exception as e:
                        # Keep the stream alive; don't push a broken event to the client
                        print(f"[SSE] snapshot error: {e}", flush=True)
                        yield ": keepalive\n\n"

                except asyncio.TimeoutError:
                    # Keepalive — SSE comment, ignored by browser
                    yield ": keepalive\n\n"

        finally:
            # Always clean up when client disconnects
            _sse_clients.discard(queue)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control":     "no-cache",
            "X-Accel-Buffering": "no",
        },
    )